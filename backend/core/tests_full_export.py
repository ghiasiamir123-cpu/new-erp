"""خروجی اکسل کامل: فقط مدیر، همهٔ جدول‌ها، بدون رمز و عکس."""

import io

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from .full_export import SYSTEM_TITLE
from .models import Employee, Project, User


class FullExportTests(TestCase):
    def setUp(self):
        self.manager = User.objects.create_user(
            username="boss", password="secret-pass", name="مدیر", role="manager",
            photo="data:image/png;base64,AAAA")
        self.clerk = User.objects.create_user(
            username="clerk", password="secret-pass", name="ثبت", role="data_entry",
            access=["dashboard", "dashboard.backup"])
        Project.objects.create(name="ویلای آرامش", code="1405-001")
        Employee.objects.create(name="علی")
        self.client = APIClient()

    def _get(self, user):
        self.client.force_authenticate(user)
        return self.client.get("/api/export/full/")

    def test_only_the_manager_role_gets_it(self):
        # تیکِ «بک‌اپ» برای نقش غیرمدیر کافی نیست.
        self.assertEqual(self._get(self.clerk).status_code, 403)
        self.client.force_authenticate(None)
        self.assertIn(self.client.get("/api/export/full/").status_code, (401, 403))

    def test_every_table_is_a_sheet_with_the_system_name_on_the_cover(self):
        response = self._get(self.manager)
        self.assertEqual(response.status_code, 200)
        self.assertIn("diwaj-erp-", response["Content-Disposition"])
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames[0], "فهرست")
        self.assertEqual(wb["فهرست"]["A1"].value, SYSTEM_TITLE)
        for sheet in ("کاربران", "پروژه‌ها", "کارگران", "حواله‌ها", "ردیف‌های حقوق", "اقلام انبار"):
            self.assertIn(sheet, wb.sheetnames)
        projects = wb["پروژه‌ها"]
        self.assertEqual(projects["A1"].value, "شناسه")
        self.assertIn("ویلای آرامش", [c.value for c in projects[2]])

    def test_passwords_and_photos_never_leave(self):
        wb = load_workbook(io.BytesIO(self._get(self.manager).content))
        users = wb["کاربران"]
        headers = [c.value for c in users[1]]
        self.assertNotIn("password", headers)
        self.assertNotIn("photo", headers)
        flat = " ".join(str(c.value) for row in users.iter_rows() for c in row)
        self.assertNotIn("pbkdf2", flat)
        self.assertNotIn("base64", flat)
