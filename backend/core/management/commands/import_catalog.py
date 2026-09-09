"""وارد کردن کاتالوگ کالا و موجودی از برگهٔ انبارگردانی.

همین یک دستور دو کار می‌کند:

  ۱. بار اول — کاتالوگ را می‌سازد (محصول، بسته، انبار، ضریب تبدیل).
  ۲. بعد از انبارگردانی — همان فایل با ستون «تعداد موجود» پرشده را می‌دهید و
     موجودی تنظیم می‌شود.

تنظیم موجودی به‌صورت «اختلاف» ثبت می‌شود: مقدار شمرده‌شده منهای موجودی فعلی.
پس اگر فایل را دوباره وارد کنید، اختلاف صفر می‌شود و چیزی دوبار حساب نمی‌شود.

    python manage.py import_catalog diwaj-stocktake-v2.xlsx --dry-run
    python manage.py import_catalog diwaj-stocktake-v2.xlsx
"""

import re
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Sum
from django.utils import timezone

from core.units import guess_units
from core.models import (
    PackConversion,
    Product,
    Sku,
    StockItem,
    StockMovement,
    Warehouse,
)

User = get_user_model()

# دسته‌هایی که کالای شیمیایی‌اند و تاریخ انقضا/بچ دارند.
BATCH_CATEGORIES = {
    "آماده‌سازی و آستر", "رویه", "هاردنر و تینر", "هاردنر و حلال",
    "پوشش نهایی و لاک", "پوشش نهایی", "رنگ، لعاب و پتینه", "روغن و واکس",
    "پیگمنت", "اکسید", "آستر", "محافظ و افزودنی", "شوینده و نگهداری",
    "پایه‌های تزئینی", "پلاستر ونیزی", "رنگ‌ها", "ریوورده — رنگ‌های تزئینی",
    "سیستم کف", "خط اورسو",
}

# نام‌هایی که کالای آتش‌زا را نشان می‌دهند.
HAZARD_WORDS = ("تینر", "حلال", "الکل", "استون", "thinner", "solvent")

COLUMNS = {
    "pkg": "شناسه بسته",
    "brand": "برند",
    "code": "کد محصول",
    "name": "نام محصول",
    "category": "دسته",
    "pack": "اندازه بسته",
    "grit": "شماره سنباده",
    "shade": "بیس / شید",
    "price": "قیمت (ریال)",
    "qty": "تعداد موجود",
    "shelf": "کد قفسه",
    "min": "حداقل موجودی",
}


def clean(v):
    return "" if v is None else str(v).strip()


def to_decimal(v):
    if v in (None, ""):
        return None
    try:
        return Decimal(str(v).replace(",", "").strip())
    except Exception:
        return None


def pack_units(pack_size):
    """«جعبه ۱۰۰ عددی» → ۱۰۰ ؛ «عدد» → None"""
    m = re.search(r"(\d+)\s*عددی", str(pack_size))
    return int(m.group(1)) if m else None


class Command(BaseCommand):
    help = "وارد کردن کاتالوگ و موجودی از فایل انبارگردانی اکسل"

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--dry-run", action="store_true",
                            help="فقط گزارش بده، چیزی ذخیره نکن")
        parser.add_argument("--user", default=None,
                            help="نام کاربری ثبت‌کنندهٔ گردش موجودی (پیش‌فرض: اولین مدیر)")
        parser.add_argument("--catalog-only", action="store_true",
                            help="فقط کاتالوگ: انبار و موجودی از این فایل خوانده نشود")

    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl نصب نیست:  pip install openpyxl")

        wb = openpyxl.load_workbook(opts["path"], data_only=True)
        sheets = [s for s in wb.sheetnames if s.startswith("انبار")]
        if not sheets:
            raise CommandError("هیچ برگه‌ای که با «انبار» شروع شود پیدا نشد.")

        actor = (User.objects.filter(username=opts["user"]).first() if opts["user"]
                 else User.objects.filter(role="manager").order_by("id").first())
        if actor is None:
            raise CommandError("کاربری برای ثبت گردش پیدا نشد (--user بدهید).")

        dry = opts["dry_run"]
        stats = {
            "products": 0, "skus": 0, "updated": 0, "warehouses": 0,
            "stock_items": 0, "counted": 0, "adjusted": 0, "conversions": 0, "skipped": 0,
        }

        catalog_only = opts["catalog_only"]
        with transaction.atomic():
            for sheet in sheets:
                self.stdout.write(self.style.MIGRATE_HEADING(f"\n=== {sheet} ==="))
                self._import_sheet(wb[sheet], sheet, actor, dry, stats, catalog_only)

            self._build_conversions(dry, stats)

            if dry:
                self.stdout.write(self.style.WARNING("\n« اجرای آزمایشی — چیزی ذخیره نشد »"))
                transaction.set_rollback(True)

        self.stdout.write("")
        for k, label in (
            ("products", "محصول ساخته‌شده"), ("skus", "بستهٔ ساخته‌شده"),
            ("updated", "بستهٔ به‌روزشده"), ("warehouses", "انبار ساخته‌شده"),
            ("stock_items", "ردیف انبار"), ("counted", "ردیف شمرده‌شده"),
            ("adjusted", "اصلاح موجودی ثبت‌شده"), ("conversions", "ضریب تبدیل"),
            ("skipped", "ردیف نادیده‌گرفته‌شده"),
        ):
            self.stdout.write(f"  {label}: {stats[k]}")

    # ---------- برگهٔ یک انبار ----------
    def _import_sheet(self, ws, sheet_name, actor, dry, stats, catalog_only=False):
        header = [clean(c.value) for c in ws[1]]
        idx = {}
        for key, title in COLUMNS.items():
            if title in header:
                idx[key] = header.index(title)
        missing = [COLUMNS[k] for k in ("pkg", "name", "pack") if k not in idx]
        if missing:
            raise CommandError(f"ستون‌های لازم در «{sheet_name}» نیست: {missing}")

        # نام برگه یک انبار واقعی نیست، فقط برگهٔ فایل سایت است. در حالت
        # «فقط کاتالوگ» انباری ساخته نمی‌شود و موجودی هم از اینجا نمی‌آید.
        warehouse = None
        if not catalog_only:
            warehouse, created = Warehouse.objects.get_or_create(
                name=sheet_name,
                defaults={"supplies_workshop": "اصفهان" in sheet_name},
            )
            if created:
                stats["warehouses"] += 1

        today = timezone.localdate()

        for row in ws.iter_rows(min_row=2, values_only=True):
            get = lambda k: clean(row[idx[k]]) if k in idx and idx[k] < len(row) else ""
            pkg = get("pkg")
            if not pkg:
                stats["skipped"] += 1
                continue

            name, code = get("name"), get("code")
            brand, category = get("brand"), get("category")

            # محصول با ترکیب کد+نام+برند شناخته می‌شود؛ واریانت‌ها زیر همان محصول
            # می‌نشینند و در سطح بسته با گرید/شید از هم جدا می‌شوند.
            product, p_created = Product.objects.get_or_create(
                code=code, name=name, brand=brand,
                defaults={
                    "category": category,
                    "batch_tracked": category in BATCH_CATEGORIES,
                    "hazardous": any(w in name.lower() for w in HAZARD_WORDS),
                    "sellable": True,
                },
            )
            if p_created:
                stats["products"] += 1

            price = to_decimal(get("price")) or Decimal(0)
            base_unit, alt_unit, alt_rate = guess_units(get("pack"))
            sku, s_created = Sku.objects.get_or_create(
                site_package_id=pkg,
                defaults={
                    "product": product, "pack_size": get("pack"),
                    "grit": get("grit"), "shade": get("shade"), "sale_price": price,
                    "base_unit": base_unit, "alt_unit": alt_unit, "alt_to_base": alt_rate,
                },
            )
            if s_created:
                stats["skus"] += 1
            else:
                # قیمت و مشخصات از سایت به‌روز می‌شود؛ قیمت خرید دست‌نخورده می‌ماند.
                sku.product = product
                sku.pack_size = get("pack")
                sku.grit = get("grit")
                sku.shade = get("shade")
                sku.sale_price = price
                fields = ["product", "pack_size", "grit", "shade", "sale_price"]
                # واحدها فقط وقتی پر می‌شوند که هنوز تعیین نشده باشند؛ تنظیم دستی
                # کاربر با هر بار وارد کردن فایل پاک نمی‌شود.
                if not sku.base_unit:
                    sku.base_unit = base_unit
                    sku.alt_unit = alt_unit
                    sku.alt_to_base = alt_rate
                    fields += ["base_unit", "alt_unit", "alt_to_base"]
                sku.save(update_fields=fields)
                stats["updated"] += 1

            if catalog_only:
                continue

            item, i_created = StockItem.objects.get_or_create(sku=sku, warehouse=warehouse)
            if i_created:
                stats["stock_items"] += 1
            shelf, min_qty = get("shelf"), to_decimal(get("min"))
            changed = []
            if shelf and item.shelf_code != shelf:
                item.shelf_code = shelf
                changed.append("shelf_code")
            if min_qty is not None and item.min_qty != min_qty:
                item.min_qty = min_qty
                changed.append("min_qty")
            if changed:
                item.save(update_fields=changed)

            counted = to_decimal(get("qty"))
            if counted is not None:
                stats["counted"] += 1
                self._apply_count(sku, warehouse, counted, actor, today, dry, stats)

    # ---------- تنظیم موجودی بر اساس شمارش ----------
    def _apply_count(self, sku, warehouse, counted, actor, today, dry, stats):
        """اختلاف شمارش با موجودی فعلی را ثبت می‌کند؛ تکرارِ بی‌ضرر."""
        on_hand = (StockMovement.objects
                   .filter(sku=sku, warehouse=warehouse)
                   .aggregate(s=Sum("qty"))["s"] or Decimal(0))
        delta = counted - on_hand
        if delta == 0:
            return
        stats["adjusted"] += 1
        if dry:
            return
        StockMovement.objects.create(
            sku=sku, warehouse=warehouse, kind=StockMovement.Kind.COUNT,
            qty=delta, date=today, ref="انبارگردانی",
            note=f"شمارش {counted} — موجودی پیشین {on_hand}",
            created_by=actor, created_by_name=actor.name or actor.username,
        )
        StockItem.objects.filter(sku=sku, warehouse=warehouse).update(counted_at=today)

    # ---------- ضریب تبدیل جعبه ↔ تکی ----------
    def _build_conversions(self, dry, stats):
        """جعبه و تکیِ یک کالا را جفت می‌کند تا شکستن بسته ممکن شود."""
        groups = {}
        for sku in Sku.objects.select_related("product"):
            key = (sku.product_id, sku.grit, sku.shade)
            groups.setdefault(key, []).append(sku)

        for skus in groups.values():
            single = next((s for s in skus if s.pack_size.strip() == "عدد"), None)
            if single is None:
                continue
            for s in skus:
                units = pack_units(s.pack_size)
                if not units or s.id == single.id:
                    continue
                if dry:
                    stats["conversions"] += 1
                    continue
                _, created = PackConversion.objects.get_or_create(
                    box_sku=s, unit_sku=single, defaults={"factor": units}
                )
                if created:
                    stats["conversions"] += 1
