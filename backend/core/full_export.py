"""خروجی اکسل کامل سامانه — همهٔ جدول‌ها، فقط برای مدیر.

خروجی قبلی در مرورگر ساخته می‌شد و فقط همان چیزی را داشت که صفحهٔ داشبورد بارگذاری
کرده بود: گزارش‌ها، پروژه‌ها، مواد و کاربران. انبار، حواله‌ها، مالی، اموال، حقوق و
راننده بیرون می‌ماند. این یکی روی سرور از روی خودِ مدل‌ها ساخته می‌شود، پس هر جدولی
که به سامانه اضافه شود خودبه‌خود در خروجی می‌آید.

هر جدول یک برگه است با ستون‌های فارسی. تاریخ‌ها شمسی، بله/خیر به‌جای True/False،
گزینه‌ها با برچسب فارسی‌شان و ارجاع‌ها با نامِ چیزی که به آن اشاره می‌کنند. برگهٔ
اول فهرست است: نام سامانه، زمان تهیه، تهیه‌کننده و تعداد ردیف هر برگه با پیوند به آن.

رمز عبور و عکس پروفایل و پیوست‌های گفتگو هرگز نوشته نمی‌شوند.
"""

from __future__ import annotations

import datetime
import decimal
import io
import json

from django.apps import apps
from django.db import models
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rest_framework.permissions import BasePermission, IsAuthenticated
from rest_framework.views import APIView

from .jalali import gregorian_to_jalali

SYSTEM_NAME = "Diwaj ERP"
SYSTEM_NAME_FA = "برنامه‌ریزی منابع سازمان"
SYSTEM_TITLE = f"{SYSTEM_NAME} ({SYSTEM_NAME_FA})"

# ستون‌هایی که هیچ‌وقت در فایل نمی‌آیند: رمز، عکس base64 و محتوای پیوست.
NEVER = {"password", "photo", "attachment"}

# ترتیب برگه‌ها و نام فارسی هر جدول. جدولی که اینجا نیست ته فایل با نام خودش می‌آید.
SHEETS = [
    ("User", "کاربران"),
    ("UserAuditLog", "تاریخچه کاربران"),
    ("Employee", "کارگران"),
    ("Project", "پروژه‌ها"),
    ("ProjectStage", "مراحل پروژه"),
    ("WorkStage", "فهرست مراحل تولید"),
    ("ProductionSettings", "تنظیمات تولید"),
    ("DailyReport", "گزارش‌های روزانه"),
    ("ReportItem", "ردیف‌های گزارش"),
    ("ReportProgress", "متراژ روزانه"),
    ("Feedback", "بازخورد گزارش‌ها"),
    ("Material", "مواد"),
    ("MaterialUsageReport", "گزارش مصرف مواد"),
    ("MaterialUsage", "ردیف‌های مصرف مواد"),
    ("MaterialUsageFeedback", "بازخورد مصرف مواد"),
    ("Driver", "رانندگان"),
    ("DriverReport", "گزارش راننده"),
    ("DriverTask", "مأموریت‌های راننده"),
    ("DriverDelay", "تأخیرهای راننده"),
    ("DriverFeedback", "بازخورد راننده"),
    ("PayrollSettings", "تنظیمات حقوق"),
    ("PayrollStaff", "پرسنل حقوق"),
    ("PayrollMonth", "ماه‌های حقوق"),
    ("PayrollEntry", "ردیف‌های حقوق"),
    ("Warehouse", "انبارها"),
    ("Location", "محل‌ها"),
    ("Supplier", "تأمین‌کنندگان"),
    ("Product", "کالاها"),
    ("Sku", "اقلام انبار"),
    ("SitePackLink", "اتصال به سایت"),
    ("PackConversion", "تبدیل بسته‌بندی"),
    ("StockItem", "موجودی اقلام"),
    ("StockBatch", "بچ‌ها"),
    ("StockMovement", "گردش انبار"),
    ("StockVoucher", "حواله‌ها"),
    ("StockVoucherLine", "ردیف‌های حواله"),
    ("StockCount", "انبارگردانی"),
    ("StockCountLine", "ردیف‌های انبارگردانی"),
    ("FamilyReview", "بازبینی خانواده‌ها"),
    ("ShopPriceSync", "همگام‌سازی قیمت سایت"),
    ("AssetInspection", "بازرسی اموال"),
    ("AssetInspectionLine", "ردیف‌های بازرسی"),
    ("AssetEvent", "رویدادهای اموال"),
    ("MaintenanceAlert", "اخطارهای نگهداری"),
    ("Conversation", "گفتگوها"),
    ("ConversationMember", "اعضای گفتگو"),
    ("Message", "پیام‌ها"),
]

FIELD_FA = {
    "id": "شناسه", "name": "نام", "code": "کد", "title": "عنوان", "note": "یادداشت",
    "desc": "شرح", "description": "شرح", "text": "متن", "status": "وضعیت",
    "active": "فعال", "is_active": "فعال", "order": "ترتیب", "date": "تاریخ",
    "created_at": "زمان ایجاد", "updated_at": "آخرین تغییر", "created_by": "ایجادکننده",
    "created_by_name": "نام ایجادکننده", "at": "زمان", "kind": "نوع", "key": "کلید",
    "level": "سطح", "label": "برچسب", "number": "شماره", "ref": "مرجع", "source": "منبع",
    "message": "پیام", "reason": "دلیل", "changes": "تغییرات", "action": "اقدام",
    # کاربر
    "username": "نام کاربری", "first_name": "نام", "last_name": "نام خانوادگی",
    "email": "ایمیل", "role": "نقش", "position": "سمت", "access": "دسترسی‌ها",
    "is_staff": "کارمند ادمین", "is_superuser": "مدیر کل", "last_login": "آخرین ورود",
    "date_joined": "تاریخ عضویت", "must_change_password": "باید رمز عوض کند",
    "target": "کاربر هدف", "target_username": "نام کاربری هدف", "actor": "انجام‌دهنده",
    "actor_name": "نام انجام‌دهنده", "user": "کاربر",
    # پروژه و تولید
    "project": "پروژه", "project_name": "نام پروژه", "owner_name": "مالک / مشتری",
    "price": "مبلغ قرارداد", "start_date": "تاریخ شروع", "due_date": "تاریخ تحویل",
    "no_area": "بدون متراژ", "general": "کار عمومی کارگاه", "base_area": "متراژ پایه",
    "closed_at": "تاریخ بستن", "close_reason": "دلیل بستن", "closed_by_name": "بسته‌شده توسط",
    "close_note": "یادداشت بستن", "closed_remaining": "متراژ باقیمانده هنگام بستن",
    "stage": "مرحله", "coefficient": "ضریب", "area": "متراژ", "done": "انجام‌شده",
    "needs_area": "نیاز به متراژ", "default_coefficient": "ضریب پیش‌فرض", "weight": "وزن",
    "importance": "اهمیت", "time_weight": "وزن زمانی",
    "labour_cost_per_hour": "هزینهٔ هر ساعت کار", "material_cost_per_m2": "هزینهٔ مواد هر متر",
    "margin_percent": "درصد سود", "hourly_cost": "هزینهٔ هر ساعت",
    # گزارش روزانه
    "shift": "شیفت", "supervisor": "سرپرست", "supervisor_name": "نام سرپرست",
    "resubmitted": "دوباره ارسال شده", "problems": "مشکلات", "report": "گزارش",
    "employee": "کارگر", "activity": "فعالیت", "hours": "ساعت", "percent": "درصد زمان",
    "manager_name": "نام مدیر",
    # مواد
    "unit": "واحد", "sku": "قلم انبار", "material": "ماده", "material_name": "نام ماده",
    "material_code": "کد ماده", "quantity": "مقدار", "recorded_by": "ثبت‌کننده",
    "recorded_by_name": "نام ثبت‌کننده", "affects_stock": "اثر روی موجودی",
    # راننده
    "driver": "راننده", "driver_name": "نام راننده",
    "morning_scheduled_time": "ساعت برنامهٔ صبح", "morning_arrival_time": "ساعت رسیدن صبح",
    "morning_passengers": "مسافران صبح", "evening_scheduled_time": "ساعت برنامهٔ عصر",
    "evening_arrival_time": "ساعت رسیدن عصر", "evening_passengers": "مسافران عصر",
    "odometer_start": "کیلومتر شروع", "odometer_end": "کیلومتر پایان", "period": "نوبت",
    "time": "ساعت", "destination": "مقصد",
    # حقوق
    "daily_hours": "ساعت کار روزانه", "ot_mult": "ضریب اضافه‌کار", "ins_rate": "نرخ بیمه",
    "tax_exempt": "معافیت مالیاتی", "components": "اجزای حقوق", "brackets": "پله‌های مالیات",
    "dept": "واحد", "married": "متأهل", "children": "تعداد فرزند", "month": "ماه",
    "staff": "پرسنل", "staff_name": "نام پرسنل", "absent_days": "روز غیبت",
    "worked_days": "روز کارکرد", "ot_hours": "ساعت اضافه‌کار", "short_hours": "ساعت کسرکار",
    "kpi": "پاداش عملکرد", "seniority": "پایه سنوات", "transport": "ایاب و ذهاب",
    "responsibility": "حق مسئولیت", "insurance_manual": "بیمهٔ دستی", "advance": "مساعده",
    "reserve": "ذخیره", "loan": "وام",
    # انبار
    "warehouse": "انبار", "to_warehouse": "انبار مقصد", "location": "محل",
    "location_name": "نام محل", "supplier": "تأمین‌کننده", "lead_time_days": "زمان تأمین (روز)",
    "supplies_workshop": "تأمین کارگاه", "product": "کالا", "brand": "برند", "category": "دسته",
    "sepidar_item_id": "کد سپیدار", "sellable": "قابل فروش", "batch_tracked": "ردیابی بچ",
    "hazardous": "خطرناک", "site_package_id": "شناسهٔ بستهٔ سایت", "warehouse_name": "نام در انبار",
    "site_name": "نام در سایت", "shop_pack_id": "شناسهٔ بسته در فروشگاه", "site_parent": "محصول سایت",
    "variant_label": "برچسب واریانت", "warehouse_code": "کد انبار", "needs_review": "نیاز به بازبینی",
    "pack_size": "اندازهٔ بسته", "base_unit": "بسته‌بندی اصلی", "alt_unit": "بسته‌بندی فرعی",
    "alt_to_base": "نرخ تبدیل", "grit": "گرید", "shade": "رنگ", "barcode": "بارکد",
    "weight_kg": "وزن (کیلو)", "sale_price": "قیمت فروش", "cost_price": "قیمت خرید",
    "site_pack": "بستهٔ سایت", "per_pack": "تعداد در بسته", "box_sku": "قلم جعبه",
    "unit_sku": "قلم تکی", "factor": "ضریب تبدیل", "batch": "بچ", "batch_no": "شمارهٔ بچ",
    "produced_on": "تاریخ تولید", "expires_on": "تاریخ انقضا", "shelf_code": "کد قفسه",
    "min_qty": "حداقل موجودی", "reserved_qty": "رزرو", "counted_at": "زمان شمارش",
    "qty": "مقدار", "entered_qty": "مقدار واردشده", "entered_unit": "واحد واردشده",
    "unit_cost": "بهای واحد", "voucher": "حواله", "usage_report": "گزارش مصرف",
    "movement_kind": "نوع حواله", "counterparty": "طرف حساب", "posted_at": "زمان ثبت نهایی",
    "posted_by_name": "ثبت نهایی توسط", "invoice_qty": "مقدار فاکتور", "unit_price": "قیمت واحد",
    "count": "انبارگردانی", "system_qty": "موجودی دفتر", "counted_qty": "موجودی شمارش‌شده",
    "posted_diff": "اختلاف ثبت‌شده", "fingerprint": "اثر انگشت", "item_count": "تعداد اقلام",
    "reviewed_by": "بازبین", "reviewed_by_name": "نام بازبین", "reviewed_at": "زمان بازبینی",
    "started_at": "شروع", "finished_at": "پایان", "ok": "موفق", "items": "اقلام",
    "updated": "به‌روزشده", "triggered_by": "اجراکننده", "triggered_by_name": "نام اجراکننده",
    # مالی
    "finance_status": "وضعیت مالی", "invoice_no": "شمارهٔ فاکتور", "invoice_date": "تاریخ فاکتور",
    "invoice_total": "جمع فاکتور", "invoice_discount": "تخفیف فاکتور", "invoice_tax": "مالیات فاکتور",
    "finance_note": "یادداشت مالی", "warehouse_reply": "پاسخ انبار", "finance_by": "تأییدکنندهٔ مالی",
    "finance_by_name": "نام تأییدکنندهٔ مالی", "finance_at": "زمان تأیید مالی",
    # اموال و نگهداری
    "is_asset": "اموال", "asset_code": "کد اموال", "holder_name": "تحویل‌گیرنده",
    "handed_over_on": "تاریخ تحویل", "asset_status": "وضعیت اموال", "asset_serial": "سریال",
    "asset_model": "مدل", "asset_supplier": "فروشندهٔ اموال", "purchase_date": "تاریخ خرید",
    "purchase_price": "قیمت خرید اموال", "warranty_until": "گارانتی تا", "asset_note": "یادداشت اموال",
    "service_interval_days": "فاصلهٔ سرویس (روز)", "useful_life_years": "عمر مفید (سال)",
    "salvage_value": "ارزش اسقاط", "inspection": "بازرسی", "present": "موجود",
    "needs_action": "نیاز به اقدام", "cost": "هزینه", "detail": "جزئیات",
    "auto_closed": "بسته‌شدن خودکار",
    # گفتگو
    "conversation": "گفتگو", "last_message_at": "آخرین پیام", "last_read_at": "آخرین خواندن",
    "joined_at": "زمان پیوستن", "left_at": "زمان خروج", "sender": "فرستنده",
    "sender_name": "نام فرستنده", "attachment_name": "نام پیوست", "attachment_kind": "نوع پیوست",
}

# سرخط برگه‌ها: سبز تیرهٔ دیواژ.
HEAD_FILL = PatternFill("solid", fgColor="1F4E3D")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="D5DDD8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MAX_CELL = 32000  # سقف اکسل ۳۲۷۶۷ نویسه است.


class IsManager(BasePermission):
    message = "خروجی اکسل کامل فقط برای مدیر است."

    def has_permission(self, request, view):
        user = request.user
        return bool(user and user.is_authenticated and user.role == "manager")


def _jalali(value):
    if isinstance(value, datetime.datetime):
        local = timezone.localtime(value) if timezone.is_aware(value) else value
        jy, jm, jd = gregorian_to_jalali(local.year, local.month, local.day)
        return f"{jy:04d}/{jm:02d}/{jd:02d} {local:%H:%M}"
    jy, jm, jd = gregorian_to_jalali(value.year, value.month, value.day)
    return f"{jy:04d}/{jm:02d}/{jd:02d}"


def _cell(obj, field):
    """مقدار یک ستون، به شکلی که در اکسل خوانا باشد."""
    if field.choices:
        return getattr(obj, f"get_{field.name}_display")() or ""
    if isinstance(field, models.ForeignKey):
        related = getattr(obj, field.name, None)
        return str(related) if related is not None else ""
    value = getattr(obj, field.attname)
    if value is None:
        return ""
    if isinstance(value, bool):
        return "بله" if value else "خیر"
    if isinstance(value, (datetime.datetime, datetime.date)):
        return _jalali(value)
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M")
    if isinstance(value, decimal.Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and len(value) > MAX_CELL:
        return value[:MAX_CELL] + "…"
    return value


def _style_header(ws, columns):
    for i, _ in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def _write_model(wb, model, sheet_name):
    fields = [f for f in model._meta.concrete_fields if f.name not in NEVER]
    ws = wb.create_sheet(sheet_name[:31])
    ws.sheet_view.rightToLeft = True
    headers = [FIELD_FA.get(f.name, f.name) for f in fields]
    ws.append(headers)
    _style_header(ws, headers)

    fks = [f.name for f in fields if isinstance(f, models.ForeignKey)]
    rows = 0
    widths = [len(h) + 2 for h in headers]
    for obj in model.objects.select_related(*fks).order_by("pk").iterator(chunk_size=2000):
        values = [_cell(obj, f) for f in fields]
        ws.append(values)
        rows += 1
        for i, v in enumerate(values):
            widths[i] = max(widths[i], min(len(str(v)), 50) + 2)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w, 8), 52)
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{rows + 1}"
    return ws.title, rows


def build_workbook(user) -> bytes:
    wb = Workbook()
    cover = wb.active
    cover.title = "فهرست"
    cover.sheet_view.rightToLeft = True

    core = apps.get_app_config("core")
    by_name = {m.__name__: m for m in core.get_models()}
    order = [(name, fa) for name, fa in SHEETS if name in by_name]
    listed = {name for name, _ in order}
    order += [(name, name) for name in sorted(by_name) if name not in listed]

    written = []
    for name, fa in order:
        written.append((fa, *_write_model(wb, by_name[name], fa)))

    now = timezone.now()
    cover["A1"] = SYSTEM_TITLE
    cover["A1"].font = Font(bold=True, size=16, color="1F4E3D")
    cover["A2"] = "گزارش کامل همهٔ اطلاعات سامانه"
    cover["A2"].font = Font(size=11, color="555555")
    cover["A4"], cover["B4"] = "زمان تهیه", _jalali(now)
    cover["A5"], cover["B5"] = "تهیه‌کننده", f"{user.name or user.username} ({user.username})"
    cover["A6"], cover["B6"] = "تعداد برگه‌ها", len(written)
    cover["A7"], cover["B7"] = "جمع ردیف‌ها", sum(r for _fa, _t, r in written)
    for row in range(4, 8):
        cover.cell(row=row, column=1).font = Font(bold=True)

    cover["A9"], cover["B9"] = "برگه", "تعداد ردیف"
    _style_header_row(cover, 9)
    for i, (fa, title, rows) in enumerate(written, start=10):
        link = cover.cell(row=i, column=1, value=fa)
        link.hyperlink = f"#'{title}'!A1"
        link.font = Font(color="1F4E3D", underline="single")
        cover.cell(row=i, column=2, value=rows)
        for col in (1, 2):
            cover.cell(row=i, column=col).border = BOX
    cover.cell(row=len(written) + 11, column=1,
               value="رمز عبور، عکس پروفایل و فایل پیوست‌ها در این خروجی نیست.").font = Font(
        size=9, color="888888")
    cover.column_dimensions["A"].width = 34
    cover.column_dimensions["B"].width = 30

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _style_header_row(ws, row):
    for col in (1, 2):
        cell = ws.cell(row=row, column=col)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BOX
        cell.alignment = Alignment(horizontal="center")


class FullExportView(APIView):
    permission_classes = [IsAuthenticated, IsManager]

    def get(self, request):
        data = build_workbook(request.user)
        now = timezone.localtime()
        jy, jm, jd = gregorian_to_jalali(now.year, now.month, now.day)
        filename = f"diwaj-erp-{jy:04d}-{jm:02d}-{jd:02d}.xlsx"
        response = HttpResponse(
            data,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response["Cache-Control"] = "no-store"
        return response
